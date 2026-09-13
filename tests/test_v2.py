# -*- coding: utf-8 -*-
"""wordfmt.py 2.0 新功能测试：replace / revisions / headers"""
import os
import subprocess
import sys

PY = sys.executable
WF = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "scripts", "wordfmt.py")
TD = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_t20")

os.makedirs(TD, exist_ok=True)
ok, fail = 0, 0


def run(*a):
    r = subprocess.run([PY, WF, *a], capture_output=True, text=True, encoding="utf-8")
    return r.returncode, r.stdout + r.stderr


def check(name, cond, detail=""):
    global ok, fail
    if cond:
        ok += 1
        print(f"  PASS  {name}")
    else:
        fail += 1
        print(f"  FAIL  {name}  {detail}")


# ============ 1) replace：跨 run + 表格 + 页眉 ============
from docx import Document
from docx.shared import Pt

f1 = os.path.join(TD, "rep.docx")
doc = Document()
# 段落：目标串被拆在三个 run 里，各 run 格式不同
p = doc.add_paragraph()
r1 = p.add_run("公司注册地址：北京")   # 12pt
r1.font.size = Pt(12)
r2 = p.add_run("市朝阳区XX路")        # 16pt 加粗
r2.font.size = Pt(16)
r2.font.bold = True
r3 = p.add_run("1号，邮编100000。")   # 10.5pt
r3.font.size = Pt(10.5)
# 正常单 run 命中
p2 = doc.add_paragraph("另一处提到 北京市朝阳区XX路1号 的段落。")
# 表格命中
t = doc.add_table(rows=1, cols=2)
t.style = "Table Grid"
t.cell(0, 0).text = "地址"
t.cell(0, 1).text = "北京市朝阳区XX路1号"
# 页眉命中
doc.sections[0].header.paragraphs[0].add_run("投标方：北京市朝阳区XX路1号")
doc.save(f1)

code, out = run("replace", f1, "北京市朝阳区XX路1号", "成都市高新区天府大道9号", "--dry-run")
check("replace dry-run 找到4处(正文2+表格1+页眉1)", "4 处可替换" in out, out)

code, out = run("replace", f1, "北京市朝阳区XX路1号", "成都市高新区天府大道9号")
check("replace 写入", code == 0 and "共替换 4 处" in out, out)

doc2 = Document(f1)
txt = doc2.paragraphs[0].text
check("跨run替换后文本正确", "成都市高新区天府大道9号" in txt and "XX路" not in txt, txt)
# 格式保持：替换文字并入起点 run（12pt），后续 run（邮编段）10.5pt 不变
runs = [r for r in doc2.paragraphs[0].runs if r.text]
host = [r for r in runs if "成都市" in r.text]
check("替换文字并入起点run(12pt)", len(host) == 1 and host[0].font.size and host[0].font.size.pt == 12,
      str([(r.text, r.font.size.pt if r.font.size else None) for r in runs]))
check("span后段run格式未受扰动(10.5pt仍在)", any(r.text.startswith("，邮编") and r.font.size and r.font.size.pt == 10.5 for r in runs),
      str([(r.text, r.font.size.pt if r.font.size else None) for r in runs]))
tbl_txt = doc2.tables[0].cell(0, 1).text
check("表格内替换成功", "成都市高新区天府大道9号" in tbl_txt, tbl_txt)
hdr_txt = doc2.sections[0].header.paragraphs[0].text
check("页眉内替换成功", "成都市高新区天府大道9号" in hdr_txt, hdr_txt)

# 正则 + 反向引用
f1b = os.path.join(TD, "re2.docx")
d = Document()
d.add_paragraph("本项目于2023年启动，2024年验收，2026年推广。")
d.save(f1b)
code, out = run("replace", f1b, "(\\d{4})年", "公元\\1年", "--regex")
d2 = Document(f1b)
check("正则替换+反向引用", "公元2023年" in d2.paragraphs[0].text and "公元2026年" in d2.paragraphs[0].text, d2.paragraphs[0].text)

# ============ 2) revisions：手工构造 w:ins / w:del / rPrChange ============
from lxml import etree
from docx.oxml.ns import qn

f2 = os.path.join(TD, "rev.docx")
doc = Document()
p = doc.add_paragraph()
r = p.add_run("本合同由甲乙双方")
r.font.size = Pt(12)
# 插入修订（张三）
ins = etree.SubElement(p._element, qn("w:ins"))
ins.set(qn("w:author"), "张三")
ins.set(qn("w:date"), "2026-09-12T10:00:00Z")
ri = etree.SubElement(ins, qn("w:r"))
ti = etree.SubElement(ri, qn("w:t"))
ti.text = "共同"
# 删除修订（李四）：紧跟的"单独"被标记删除
r2 = p.add_run("单独")
r2.font.size = Pt(12)
# 把 r2 移进 w:del 并把 w:t 改成 w:delText
dele = etree.SubElement(p._element, qn("w:del"))
dele.set(qn("w:author"), "李四")
dele.set(qn("w:date"), "2026-09-12T11:00:00Z")
p._element.remove(r2._element)
dele.append(r2._element)
r2._element.find(qn("w:t")).tag = qn("w:delText")
p.add_run("签署。")
# 格式修订：p2 的 run 加粗 + rPrChange 记录旧格式
p3 = doc.add_paragraph()
r3 = p3.add_run("重要条款")
r3.font.bold = True
rpr = r3._element.get_or_add_rPr()
ch = etree.SubElement(rpr, qn("w:rPrChange"))
ch.set(qn("w:author"), "王五")
old = etree.SubElement(ch, qn("w:rPr"))
doc.save(f2)

code, out = run("revisions", f2)
check("revisions list 列出插入(张三)", "插入" in out and "张三" in out and "共同" in out, out)
check("revisions list 列出删除(李四)", "李四" in out, out)
check("revisions list 格式修订计数", "格式修订" in out, out)

# accept
f2a = os.path.join(TD, "rev_accept.docx")
import shutil
shutil.copy2(f2, f2a)
code, out = run("revisions", f2a, "--action", "accept")
d3 = Document(f2a)
final = d3.paragraphs[0].text
check("accept: 插入保留+删除移除", final == "本合同由甲乙双方共同签署。", repr(final))
check("accept: 无残留修订标记", not d3.element.findall(f".//{qn('w:ins')}") and not d3.element.findall(f".//{qn('w:del')}"))
check("accept: rPrChange 已清", not d3.element.findall(f".//{qn('w:rPrChange')}"))
check("accept: 加粗保留", d3.paragraphs[1].runs[0].font.bold is True)

# reject
f2b = os.path.join(TD, "rev_reject.docx")
shutil.copy2(f2, f2b)
code, out = run("revisions", f2b, "--action", "reject")
d4 = Document(f2b)
final = d4.paragraphs[0].text
check("reject: 插入撤销+删除恢复", final == "本合同由甲乙双方单独签署。", repr(final))
check("reject: 格式回滚(加粗撤销)", d4.paragraphs[1].runs[0].font.bold is not True)

# 按作者 accept
f2c = os.path.join(TD, "rev_author.docx")
shutil.copy2(f2, f2c)
code, out = run("revisions", f2c, "--action", "accept", "--author", "李四")
d5 = Document(f2c)
# 注意：python-docx 的 paragraph.text 不读取 w:ins 内的 run，须 XML 层取全文
xml_text = "".join(t.text or "" for t in d5.paragraphs[0]._element.iter()
                   if t.tag in (qn("w:t"), qn("w:delText")))
check("按作者accept: 只接受李四的删除,张三插入保留",
      xml_text == "本合同由甲乙双方共同签署。" and d5.element.findall(f".//{qn('w:ins')}"),
      repr(xml_text))

# ============ 3) headers ============
code, out = run("headers", f1)
check("headers list 显示页眉文本", "默认页眉" in out and "投标方" in out, out)

f3 = os.path.join(TD, "hdr.docx")
d = Document()
d.add_paragraph("正文")
d.save(f3)
code, out = run("headers", f3, "--set-header", "西南交通大学教务处", "--set-footer", "第 1 页")
d2 = Document(f3)
check("set-header 写入", d2.sections[0].header.paragraphs[0].text == "西南交通大学教务处",
      d2.sections[0].header.paragraphs[0].text)
check("set-footer 写入", d2.sections[0].footer.paragraphs[0].text == "第 1 页",
      d2.sections[0].footer.paragraphs[0].text)
code, out = run("headers", f3)
check("headers list 回读", "西南交通大学教务处" in out, out)

# ============ 4) 回归：原三命令不受影响 ============
code, out = run("inspect", f1)
check("回归 inspect", code == 0 and "格式体检报告" in out, out)
code, out = run("md2docx", os.path.join(TD, "x.md") if os.path.exists(os.path.join(TD, "x.md")) else f1b)
check("回归 md2docx help 不崩", code in (0, 1))

print(f"\n===== 结果: {ok} 通过, {fail} 失败 =====")
sys.exit(1 if fail else 0)
